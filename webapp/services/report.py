# -*- coding: utf-8 -*-
"""
Выгрузка расчёта экономии в Excel.

Энергетик пересылает расчёт руководству и в бухгалтерию, где его открывают в
Excel и пересчитывают по-своему. Поэтому здесь не картинка, а таблицы с
числами без округления и единицами в заголовках, плюс почасовое расписание:
по нему можно проверить любую цифру страницы.
"""

import io
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from webapp.texts import source_ru


def _sheet(wb: Workbook, title: str, header: List[str], rows: List[List[Any]],
           first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, name in enumerate(header, start=1):
        width = max([len(str(name))] + [len(f"{r[i - 1]}") for r in rows[:200] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    return ws


def workbook(result: Dict[str, Any], passport: List[Dict[str, str]], params: Dict[str, Any],
             fin: Optional[Dict[str, Any]], whatif: Optional[List[Dict[str, Any]]]) -> bytes:
    b, bat = result["best"], result["battery"]
    wb = Workbook()
    summary = [
        ["Период оценки", f"{result['period'][0]}–{result['period'][1]}", "сут.", result["n_days"]],
        ["Ценовая категория", result["tariff"]["category"], "", ""],
        ["Ставки", "примерные" if result["tariff"]["example"] else "заданы пользователем", "", ""],
        ["Прогноз для плана заряда", source_ru(result["best_source"]), "", ""],
        ["Мощность накопителя", bat["power_kw"], "кВт", ""],
        ["Ёмкость накопителя", bat["capacity_kwh"], "кВт·ч", ""],
        ["Вложения", bat["capex_rub"], "₽", ""],
        ["Чистая экономия за период", b["net_savings"], "₽", ""],
        ["Экономия в год, оценка", b["annual_mean"], "₽/год", ""],
        ["Экономия в год, нижний край (90 %)", b["annual_lo"], "₽/год", ""],
        ["Экономия в год, верхний край (90 %)", b["annual_hi"], "₽/год", ""],
        ["Простая окупаемость", b["payback_years"], "лет", ""],
    ]
    risk = result.get("peak_risk")
    if risk:
        summary += [["Экономия в год, если час пика региона — любой плановый", risk["annual_mean"], "₽/год", ""],
                    ["Экономия в год, если час пика региона — самый неудачный", risk["annual_worst"], "₽/год", ""]]
    if fin:
        p = fin["params"]
        summary += [["Ставка дисконтирования", p["discount"], "доля", ""],
                    ["Срок службы", p["life"], "лет", ""],
                    ["Индексация тарифов", p["growth"], "доля в год", ""],
                    ["NPV, оценка", fin["mean"]["npv"], "₽", ""],
                    ["IRR, оценка", fin["mean"]["irr"], "доля", ""],
                    ["Дисконтированная окупаемость", fin["mean"]["payback"], "лет", ""]]
    _sheet(wb, "Итог", ["Показатель", "Значение", "Единица", "Сутки"], summary, first=True)

    _sheet(wb, "Паспорт", ["Уровень", "Условие"], [[p["level"], p["text"]] for p in passport])
    _sheet(wb, "Составляющие", ["Составляющая", "За период, ₽"],
           [[p["label"], p["value"]] for p in result["parts"]])
    _sheet(wb, "Прогнозы", ["Прогноз", "Ошибка, кВт·ч в час", "За период, ₽", "В год, ₽",
                            "Нижний край, ₽/год", "Верхний край, ₽/год", "От предела"],
           [[source_ru(s["source"]), s.get("forecast_MAE"), s["net_savings"], s["annual_mean"],
             s["annual_lo"], s["annual_hi"], s.get("share_of_bound")] for s in result["sources"]])
    _sheet(wb, "По месяцам", ["Месяц", "Суток с данными", "Снижение счёта, ₽", "Неполный"],
           [[m["month"], m["days"], m["gross_savings"], "да" if m["partial"] else ""]
            for m in result["monthly"]])
    if whatif:
        _sheet(wb, "Что если", ["Фактор", "Значение", "В год, ₽", "Нижний край", "Верхний край",
                                "Окупаемость, лет"],
               [[p["factor"], p["value"], p["annual"], p["annual_lo"], p["annual_hi"], p["payback"]]
                for p in whatif])
    from datetime import datetime, timezone
    s = result["schedule"]
    _sheet(wb, "Почасово", ["Время", "Нагрузка, кВт·ч", "Из сети с накопителем, кВт·ч",
                            "Заряд накопителя, кВт·ч"],
           [[datetime.fromtimestamp(t, tz=timezone.utc).replace(tzinfo=None), l, g, c]
            for t, l, g, c in zip(s["t"], s["load"], s["grid"], s["soc"])])
    _sheet(wb, "Параметры", ["Параметр", "Значение"],
           [[k, ", ".join(v) if isinstance(v, list) else v] for k, v in params.items()])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
